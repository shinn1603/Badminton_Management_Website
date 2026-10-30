# -*- coding: utf-8 -*-
"""
doc_helper.py: Các hàm hỗ trợ định dạng văn bản chuẩn Times New Roman, cỡ chữ 13, màu đen,
bảng biểu và khối mã nguồn chuyên nghiệp cho tài liệu UTE Sport.
"""

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BLACK = RGBColor(0, 0, 0)
DARK_GRAY = RGBColor(50, 50, 50)
BORDER_GRAY = "CCCCCC"
BG_LIGHT_GRAY = "F8F9FA"
BG_HEADER_GRAY = "E9ECEF"
BG_NOTE_BOX = "F0FDF4"
BORDER_NOTE_GREEN = "059669"

def init_document():
    doc = docx.Document()
    
    # 1. Cấu hình lề trang tiêu chuẩn 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # 2. Cấu hình Font mặc định: Times New Roman, cỡ 13pt, chữ đen
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = BLACK
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)
    
    return doc

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLACK
    return p

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.italic = True
    run.font.color.rgb = DARK_GRAY
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = BLACK
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = BLACK
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = BLACK
    return p

def add_paragraph(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Times New Roman'
        r_prefix.font.size = Pt(13)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = BLACK
        
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(13)
    r_text.font.color.rgb = BLACK
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Times New Roman'
        r_prefix.font.size = Pt(13)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = BLACK
        
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(13)
    r_text.font.color.rgb = BLACK
    return p

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Mã nguồn minh họa: {caption}")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(11.5)
        r_cap.font.bold = True
        r_cap.font.italic = True
        r_cap.font.color.rgb = DARK_GRAY
        
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, BG_LIGHT_GRAY)
    set_table_borders(tbl, color="D1D5DB", sz="4")
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    run = p.add_run(code_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    # Khoảng cách sau bảng
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def add_ascii_diagram(doc, diagram_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Sơ đồ quy trình: {caption}")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(12)
        r_cap.font.bold = True
        r_cap.font.color.rgb = BLACK
        
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_table_borders(tbl, color="CBD5E1", sz="6")
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    run = p.add_run(diagram_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def add_note_box(doc, text, title="LƯU Ý QUAN TRỌNG CHO NGƯỜI MỚI HỌC LẬP TRÌNH"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, BG_NOTE_BOX)
    
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{BORDER_NOTE_GREEN}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    r_title = p.add_run(f"💡 {title}:\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(12)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(4, 120, 87)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(12)
    r_text.font.color.rgb = BLACK
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def add_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="D3D3D3", sz="4")
    
    # Format Headers
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        set_cell_background(hdr_cells[i], BG_HEADER_GRAY)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = BLACK
            
    # Format Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = BLACK
                
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)
    return tbl
